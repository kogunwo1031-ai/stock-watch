"""
주간 갱신용 빌드 스크립트
입력: research.json(고정 리서치), quotes.csv(이번 주 시세), [선택] prev_quotes.csv(지난 조회 시세)
출력: 저평가_우량주_리스트_<날짜>.xlsx, data.json(대시보드용)
사용: python3 build.py 2026-09-30
"""
import csv, json, os, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

RUN_DATE = sys.argv[1] if len(sys.argv) > 1 else datetime.date.today().isoformat()
R = json.load(open("research.json", encoding="utf-8"))
Q = {r["ticker"]: r for r in csv.DictReader(open("quotes.csv", encoding="utf-8"))}
P = {}
if os.path.exists("prev_quotes.csv"):
    P = {r["ticker"]: r for r in csv.DictReader(open("prev_quotes.csv", encoding="utf-8"))}

def num(s):
    s = (s or "").strip()
    return float(s) if s else None
def mcap_b(s):
    s = (s or "").strip()
    if not s: return None
    return float(s[:-1]) * 1000 if s.endswith("T") else float(s[:-1])
RATING = {"Strong Buy": "강력 매수", "Buy": "매수", "Hold": "보유", "Sell": "매도", "Strong Sell": "강력 매도"}

SECTOR_COLOR = {  # 배지(진한색)
 "정보기술": "0D2B52", "커뮤니케이션서비스": "7030A0", "산업재": "44546A", "소재": "7F7F7F",
 "에너지": "C55A11", "에너지(신재생)": "548235", "금융": "0B6E69", "헬스케어": "C00000",
 "필수소비재": "8B6914", "임의소비재": "C2185B", "부동산(리츠)": "5D4037", "유틸리티": "1565C0"}
def light(hexc, k=0.85):
    r, g, b = (int(hexc[i:i+2], 16) for i in (0, 2, 4))
    return "".join(f"{int(c + (255 - c) * k):02X}" for c in (r, g, b))
SECTOR_ORDER = list(SECTOR_COLOR.keys())
GRADE_FILL = {"A": "C6EFCE", "B": "FFEB9C", "C": "EDEDED"}
GRADE_TEXT = {"A": "006100", "B": "7F6000", "C": "595959"}

H = json.load(open("history.json")) if os.path.exists("history.json") else {}
def iso(n): s = str(n); return f"{s[:4]}-{s[4:6]}-{s[6:]}"

rows = []
missing = []
for s in R["stocks"]:
    q = Q.get(s["t"])
    if not q: missing.append(s["t"]); continue
    p = P.get(s["t"])
    d = dict(s)
    d.update(price=float(q["price"]), qdate=q["date"], mcap=mcap_b(q["mcap"]), pe=num(q["pe"]), fpe=num(q["fpe"]),
             div=num(q["div"]), rating=RATING.get(q["rating"], q["rating"]), tgt=num(q["target"]),
             lo=num(q["lo"]), hi=num(q["hi"]), prev=float(p["price"]) if p else None, prevdate=p["date"] if p else None)
    # 일봉이 시세 페이지보다 같거나 최신이면 일봉 종가를 현재가로 사용(날짜가 명시된 확정 종가)
    h = H.get(s["t"], [])
    d["src_price"] = "시세 페이지"
    if h and iso(h[-1][0]) >= d["qdate"]:
        d["price"], d["qdate"], d["src_price"] = h[-1][4], iso(h[-1][0]), "일봉 종가"
    # 직전 거래일 종가(전일 대비 계산용)
    before = [r for r in h if iso(r[0]) < d["qdate"]]
    d["ref_close"], d["ref_date"] = (before[-1][4], iso(before[-1][0])) if before else (None, None)
    if d["lo"] is not None: d["lo"] = min(d["lo"], d["price"])
    if d["hi"] is not None: d["hi"] = max(d["hi"], d["price"])
    if d.get("note") and not d["risk"].startswith("⚠"):
        d["risk"] = "⚠ " + d["note"] + " / " + d["risk"]
    d["h"] = h
    rows.append(d)
if missing: print("시세 없음:", missing)

go = {"A": 0, "B": 1, "C": 2}
detail_rows = sorted(rows, key=lambda d: (go[d["grade"]], SECTOR_ORDER.index(d["major"]), d["t"]))
view_rows = sorted(rows, key=lambda d: (SECTOR_ORDER.index(d["major"]), go[d["grade"]], d["t"]))

F = "Arial"
def font(size=10, bold=False, color="000000", italic=False, underline=None):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic, underline=underline)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = lambda c: PatternFill("solid", fgColor=c)
HDR = "1F3864"

wb = Workbook()

# ======================= 상세 데이터 =======================
DS = "상세 데이터"
ws = wb.active; ws.title = DS
cols = [  # key, header, width, group
 ("no","No",5,"기업 정보"),("t","티커\n(클릭=회사 상세)",10,"기업 정보"),("name","회사명",24,"기업 정보"),
 ("major","대분류 섹터",14,"기업 정보"),("sub","세부 산업",20,"기업 정보"),("desc","무슨 사업을 하나",46,"기업 정보"),
 ("rev","주요 매출원·대표 제품",40,"기업 정보"),("hq","본사",12,"기업 정보"),
 ("scen","시나리오",11,"판단"),("grade","종합등급",8,"판단"),
 ("price","현재가($)",11,"시세·밸류에이션"),("qdate","시세 기준일",11,"시세·밸류에이션"),
 ("prev","지난 조회 가격($)",11,"시세·밸류에이션"),("chg","지난 조회 대비",10,"시세·밸류에이션"),
 ("mcap","시가총액($십억)",11,"시세·밸류에이션"),("lo","52주 최저($)",11,"시세·밸류에이션"),("hi","52주 최고($)",11,"시세·밸류에이션"),
 ("off52","52주 고점 대비",10,"시세·밸류에이션"),("pe","PER",9,"시세·밸류에이션"),("fpe","예상 PER",9,"시세·밸류에이션"),
 ("div","배당수익률",9,"시세·밸류에이션"),
 ("rating","애널리스트 의견",10,"애널리스트"),("tgt","애널리스트 평균 목표가($)",12,"애널리스트"),("upside","목표가 대비 상승여력",11,"애널리스트"),
 ("fv","모닝스타 공정가치($)",12,"모닝스타"),("disc","모닝스타 발표 할인율(입력)",12,"모닝스타"),("msdisc","모닝스타 기준 할인율",12,"모닝스타"),
 ("moat","해자",9,"모닝스타"),("unc","불확실성",9,"모닝스타"),
 ("val","저평가 근거",34,"근거·리스크"),("thesis","핵심 투자포인트",40,"근거·리스크"),("risk","주요 리스크",32,"근거·리스크"),
 ("mdate","모닝스타·리포트 기준일",12,"근거·리스크"),("src","출처",30,"근거·리스크"),
]
L = {k: get_column_letter(i) for i, (k, *_ ) in enumerate(cols, start=1)}
grp_color = {"기업 정보": "2E75B6", "판단": "7030A0", "시세·밸류에이션": "1F3864", "애널리스트": "375623", "모닝스타": "833C0B", "근거·리스크": "595959"}
c = 1
while c <= len(cols):
    g = cols[c-1][3]; st = c
    while c <= len(cols) and cols[c-1][3] == g: c += 1
    ws.merge_cells(start_row=1, start_column=st, end_row=1, end_column=c-1)
    ws.cell(row=1, column=st, value=g).font = font(10, True, "FFFFFF")
    ws.cell(row=1, column=st).alignment = Alignment(horizontal="center")
    for k in range(st, c): ws.cell(row=1, column=k).fill = fill(grp_color[g])
for i, (k, h, w, g) in enumerate(cols, start=1):
    cell = ws.cell(row=2, column=i, value=h)
    cell.font = font(10, True, "FFFFFF"); cell.fill = fill(grp_color[g]); cell.border = border
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[2].height = 42
R0 = 3; LAST = R0 + len(detail_rows) - 1
INPUT_KEYS = {"price","prev","mcap","lo","hi","pe","fpe","div","tgt","fv","disc"}
for n, d in enumerate(detail_rows):
    i = R0 + n
    v = dict(d)
    v["no"] = n + 1
    v["prev"] = d["prev"] if d["prev"] is not None else "첫 조회"
    v["chg"] = f'=IF(AND(ISNUMBER({L["price"]}{i}),ISNUMBER({L["prev"]}{i})),{L["price"]}{i}/{L["prev"]}{i}-1,"첫 조회")'
    v["off52"] = f'=IF(AND(ISNUMBER({L["price"]}{i}),ISNUMBER({L["hi"]}{i})),{L["price"]}{i}/{L["hi"]}{i}-1,"")'
    v["upside"] = f'=IF(AND(ISNUMBER({L["price"]}{i}),ISNUMBER({L["tgt"]}{i})),{L["tgt"]}{i}/{L["price"]}{i}-1,"목표가 없음")'
    v["msdisc"] = (f'=IF(AND(ISNUMBER({L["price"]}{i}),ISNUMBER({L["fv"]}{i})),1-{L["price"]}{i}/{L["fv"]}{i},'
                   f'IF(ISNUMBER({L["disc"]}{i}),{L["disc"]}{i},"모닝스타 수치 없음"))')
    for key, dflt in (("pe","해당없음(적자 등)"),("fpe","해당없음"),("div","무배당"),("fv","비공개"),("disc","해당 없음"),
                      ("mcap","확인 불가"),("lo","확인 불가"),("hi","확인 불가"),("tgt","목표가 없음")):
        if v.get(key) is None: v[key] = dflt
    for j, (key, *_ ) in enumerate(cols, start=1):
        cell = ws.cell(row=i, column=j, value=v[key])
        cell.border = border
        cell.font = font(10, color="0000FF") if key in INPUT_KEYS and isinstance(v[key], (int, float)) else font(10)
        cell.alignment = Alignment(vertical="top", wrap_text=key in ("name","sub","desc","rev","val","thesis","risk","src"),
                                   horizontal="center" if key in ("no","t","hq","scen","grade","qdate","rating","moat","unc","mdate") else None)
    tc = ws.cell(row=i, column=2)
    tc.hyperlink = f"https://stockanalysis.com/stocks/{d['t'].lower()}/company/"
    tc.font = font(10, True, SECTOR_COLOR[d["major"]], underline="single"); tc.fill = fill(light(SECTOR_COLOR[d["major"]], 0.85))
    ws.cell(row=i, column=cols.index(next(x for x in cols if x[0]=="major"))+1).fill = fill(light(SECTOR_COLOR[d["major"]]))
    for key in ("price","prev","lo","hi","tgt","fv"): ws[f"{L[key]}{i}"].number_format = '$#,##0.00'
    ws[f"{L['mcap']}{i}"].number_format = '#,##0.0'
    for key in ("pe","fpe"): ws[f"{L[key]}{i}"].number_format = '0.0"배"'
    for key in ("chg","off52","div","upside","disc","msdisc"): ws[f"{L[key]}{i}"].number_format = '0.0%'
ws.freeze_panes = "D3"
ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{LAST}"
gcol = L["grade"]
for g, col in GRADE_FILL.items():
    ws.conditional_formatting.add(f"{gcol}{R0}:{gcol}{LAST}", CellIsRule(operator="equal", formula=[f'"{g}"'], fill=fill(col), font=font(10, True, GRADE_TEXT[g])))
for key in ("upside", "msdisc"):
    cl = L[key]
    ws.conditional_formatting.add(f"{cl}{R0}:{cl}{LAST}", FormulaRule(formula=[f"AND(ISNUMBER({cl}{R0}),{cl}{R0}>=0.3)"], font=font(10, True, "006100"), fill=fill("E2EFDA")))
    ws.conditional_formatting.add(f"{cl}{R0}:{cl}{LAST}", FormulaRule(formula=[f"AND(ISNUMBER({cl}{R0}),{cl}{R0}<0)"], font=font(10, color="C00000")))
cl = L["chg"]
ws.conditional_formatting.add(f"{cl}{R0}:{cl}{LAST}", FormulaRule(formula=[f"AND(ISNUMBER({cl}{R0}),{cl}{R0}<0)"], font=font(10, color="C00000")))
ws.conditional_formatting.add(f"{cl}{R0}:{cl}{LAST}", FormulaRule(formula=[f"AND(ISNUMBER({cl}{R0}),{cl}{R0}>0)"], font=font(10, color="006100")))
dv = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"{gcol}{R0}:{gcol}{LAST}")

def ref(key, tick_cell):
    """상세 데이터 시트에서 티커로 값을 찾아오는 수식"""
    return (f"INDEX('{DS}'!${L[key]}${R0}:${L[key]}${LAST},"
            f"MATCH({tick_cell},'{DS}'!${L['t']}${R0}:${L['t']}${LAST},0))")

def refn(key, tick_cell):
    """숫자가 아니면 '—'로 표시"""
    x = ref(key, tick_cell)
    return f'IF(ISNUMBER({x}),{x},"—")'

# ======================= 한눈에 보기 =======================
ov = wb.create_sheet("한눈에 보기", 0)
ov_cols = [("",1.5),("티커",9),("회사명",22),("무슨 사업을 하나",52),("등급",6),("현재가($)",10),("지난 조회 대비",10),
           ("목표가 상승여력",11),("모닝스타 할인율",11),("예상 PER",8),("배당",8),("52주 고점 대비",10),("시나리오",11)]
ov["A1"] = f"저평가 우량주 한눈에 보기  ·  시세 기준 {max(d['qdate'] for d in rows)}  ·  갱신 {RUN_DATE}"
ov["A1"].font = font(14, True, HDR)
ov["A2"] = "섹터별로 묶었습니다. 왼쪽 티커 색이 섹터 색입니다. 티커를 누르면 회사 상세 페이지가 열립니다. 녹색 굵은 숫자 = 30% 이상, 빨간 숫자 = 마이너스, '—' = 해당 수치 없음(첫 갱신 포함)."
ov["A2"].font = font(9, italic=True, color="595959")
for j, (h, w) in enumerate(ov_cols, start=1):
    ov.column_dimensions[get_column_letter(j)].width = w
r = 4
for sec in SECTOR_ORDER:
    grp = [d for d in view_rows if d["major"] == sec]
    if not grp: continue
    ov.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(ov_cols))
    h = ov.cell(row=r, column=1, value=f"{sec}  ({len(grp)}개)")
    h.font = font(12, True, "FFFFFF"); h.fill = fill(SECTOR_COLOR[sec]); h.alignment = Alignment(vertical="center", indent=1)
    ov.row_dimensions[r].height = 24
    r += 1
    for j, (hname, _) in enumerate(ov_cols, start=1):
        cell = ov.cell(row=r, column=j, value=hname)
        cell.font = font(9, True, "404040"); cell.fill = fill(light(SECTOR_COLOR[sec], 0.9))
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = border
    r += 1
    for d in grp:
        tcell = f"$B{r}"
        vals = ["", d["t"], d["name"], f"={ref('desc', tcell)}", f"={ref('grade', tcell)}", f"={ref('price', tcell)}",
                f"={refn('chg', tcell)}", f"={refn('upside', tcell)}", f"={refn('msdisc', tcell)}",
                f"={ref('fpe', tcell)}", f"={ref('div', tcell)}", f"={refn('off52', tcell)}", f"={ref('scen', tcell)}"]
        for j, val in enumerate(vals, start=1):
            cell = ov.cell(row=r, column=j, value=val); cell.border = border
            cell.font = font(10); cell.alignment = Alignment(vertical="center", wrap_text=(j in (3, 4)), horizontal="center" if j in (2,5,13) or j in (7,8,9,12) else None)
        ov.cell(row=r, column=1).fill = fill(SECTOR_COLOR[sec])
        a = ov.cell(row=r, column=2)
        a.font = font(11, True, SECTOR_COLOR[sec], underline="single"); a.fill = fill(light(SECTOR_COLOR[sec], 0.85))
        a.hyperlink = f"https://stockanalysis.com/stocks/{d['t'].lower()}/company/"
        ov.cell(row=r, column=3).font = font(10, True)
        ov.cell(row=r, column=4).font = font(9, color="404040")
        ov.cell(row=r, column=6).number_format = '$#,##0.00'
        ov.cell(row=r, column=10).number_format = '0.0"배"'
        for j in (7,8,9,11,12): ov.cell(row=r, column=j).number_format = '0.0%'
        ov.row_dimensions[r].height = 30
        r += 1
    r += 1
OV_LAST = r
for g, col in GRADE_FILL.items():
    ov.conditional_formatting.add(f"E4:E{OV_LAST}", CellIsRule(operator="equal", formula=[f'"{g}"'], fill=fill(col), font=font(11, True, GRADE_TEXT[g])))
for cl in ("H", "I"):
    ov.conditional_formatting.add(f"{cl}4:{cl}{OV_LAST}", FormulaRule(formula=[f"AND(ISNUMBER({cl}4),{cl}4>=0.3)"], font=font(10, True, "006100"), fill=fill("E2EFDA")))
for cl in ("G", "H", "I", "L"):
    ov.conditional_formatting.add(f"{cl}4:{cl}{OV_LAST}", FormulaRule(formula=[f"AND(ISNUMBER({cl}4),{cl}4<0)"], font=font(10, color="C00000")))
ov.freeze_panes = "D4"
ov.sheet_view.showGridLines = False

# ======================= 종목 카드 =======================
cd = wb.create_sheet("종목 카드", 1)
cd.sheet_view.showGridLines = False
cd["A1"] = f"종목 카드  ·  시세 기준 {max(d['qdate'] for d in rows)}"; cd["A1"].font = font(14, True, HDR)
cd["A2"] = "카드 위쪽 띠 색 = 섹터 색. 숫자는 '상세 데이터' 시트에서 자동으로 가져옵니다."; cd["A2"].font = font(9, italic=True, color="595959")
PER_ROW, CW = 3, 4             # 한 줄에 카드 3장, 카드 폭 4열
for k in range(PER_ROW):
    base = 1 + k * (CW + 1)
    for j, w in enumerate((13, 13, 13, 13)):
        cd.column_dimensions[get_column_letter(base + j)].width = w
    cd.column_dimensions[get_column_letter(base + CW)].width = 3
CARD_H = [26, 16, 44, 14, 20, 14, 20, 40, 32, 8]
top = 4
for n, d in enumerate(view_rows):
    band, pos = divmod(n, PER_ROW)
    r0 = top + band * len(CARD_H)
    c0 = 1 + pos * (CW + 1)
    if pos == 0:
        for k, hgt in enumerate(CARD_H): cd.row_dimensions[r0 + k].height = hgt
    col = SECTOR_COLOR[d["major"]]; lt = light(col, 0.9)
    C = lambda rr, cc: cd.cell(row=r0 + rr, column=c0 + cc)
    def merge(rr, c1, c2): cd.merge_cells(start_row=r0+rr, start_column=c0+c1, end_row=r0+rr, end_column=c0+c2)
    # 0: 티커 배지 + 회사명
    t = C(0, 0); t.value = d["t"]; t.font = font(12, True, col, underline="single"); t.fill = fill(light(col, 0.85))
    t.alignment = Alignment(horizontal="center", vertical="center"); t.hyperlink = f"https://stockanalysis.com/stocks/{d['t'].lower()}/company/"
    merge(0, 1, 3); nm = C(0, 1); nm.value = d["name"]; nm.font = font(11, True, "FFFFFF"); nm.fill = fill(col)
    nm.alignment = Alignment(vertical="center", indent=1, shrink_to_fit=True)
    for cc in (2, 3): C(0, cc).fill = fill(col)
    tick = f"${get_column_letter(c0)}${r0}"
    # 1: 섹터 · 세부산업
    merge(1, 0, 3); s1 = C(1, 0); s1.value = f"{d['major']} · {d['sub']} · {d['hq']}"; s1.font = font(9, color="595959")
    # 2: 사업 설명
    merge(2, 0, 3); s2 = C(2, 0); s2.value = f"={ref('desc', tick)}"; s2.font = font(10); s2.alignment = Alignment(wrap_text=True, vertical="top")
    # 3-4: 현재가 | 목표가 상승여력 | 모닝스타 할인율 | 등급
    labels1 = [("현재가", "price", '$#,##0.00'), ("목표가 상승여력", "upside", '0.0%'), ("모닝스타 할인율", "msdisc", '0.0%'), ("등급", "grade", None)]
    labels2 = [("예상 PER", "fpe", '0.0"배"'), ("배당", "div", '0.0%'), ("52주 고점 대비", "off52", '0.0%'), ("지난 조회 대비", "chg", '0.0%')]
    for rr, labels in ((3, labels1), (5, labels2)):
        for cc, (lab, key, fmt) in enumerate(labels):
            lc = C(rr, cc); lc.value = lab; lc.font = font(8, color="7F7F7F"); lc.fill = fill(lt); lc.alignment = Alignment(horizontal="center")
            vc = C(rr + 1, cc); vc.value = f"={refn(key, tick)}" if key in ("msdisc", "chg", "upside", "off52") else f"={ref(key, tick)}"; vc.font = font(11, True); vc.fill = fill(lt)
            vc.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
            if fmt: vc.number_format = fmt
    # 7: 포인트, 8: 리스크
    merge(7, 0, 3); p = C(7, 0); p.value = f"=\"포인트: \"&{ref('thesis', tick)}"; p.font = font(9, color="1F3864"); p.alignment = Alignment(wrap_text=True, vertical="top")
    merge(8, 0, 3); q = C(8, 0); q.value = f"=\"리스크: \"&{ref('risk', tick)}"; q.font = font(9, color="C00000"); q.alignment = Alignment(wrap_text=True, vertical="top")
    # 카드 테두리
    edge = Side(style="thin", color=col)
    for rr in range(0, 9):
        for cc in range(0, CW):
            cell = C(rr, cc)
            cell.border = Border(left=edge if cc == 0 else None, right=edge if cc == CW-1 else None,
                                 top=edge if rr == 0 else None, bottom=edge if rr == 8 else None)
CD_LAST = top + ((len(view_rows) - 1) // PER_ROW + 1) * len(CARD_H)
for g, colr in GRADE_FILL.items():
    for k in range(PER_ROW):
        gl = get_column_letter(1 + k * (CW + 1) + 3)
        cd.conditional_formatting.add(f"{gl}4:{gl}{CD_LAST}", CellIsRule(operator="equal", formula=[f'"{g}"'], fill=fill(colr), font=font(12, True, GRADE_TEXT[g])))
for k in range(PER_ROW):
    for off in (1, 2):
        gl = get_column_letter(1 + k * (CW + 1) + off)
        cd.conditional_formatting.add(f"{gl}4:{gl}{CD_LAST}", FormulaRule(formula=[f"AND(ISNUMBER({gl}4),{gl}4>=0.3)"], font=font(11, True, "006100")))
        cd.conditional_formatting.add(f"{gl}4:{gl}{CD_LAST}", FormulaRule(formula=[f"AND(ISNUMBER({gl}4),{gl}4<0)"], font=font(11, True, "C00000")))

# ======================= 섹터 가이드 =======================
M = f"'{DS}'!"
rng = lambda key: f"{M}${L[key]}${R0}:${L[key]}${LAST}"
g = wb.create_sheet("섹터 가이드")
g.sheet_view.showGridLines = False
g["A1"] = "섹터별로 무슨 일을 하고, 금리에 따라 어떻게 움직이나"; g["A1"].font = font(13, True, HDR)
gh = ["대분류 섹터","어떤 업종인가","돈 버는 방식","금리 인상기","금리 인하기","체크할 지표","종목 수","A등급 수","평균 목표가 상승여력"]
gw = [18,40,32,34,34,32,8,8,12]
for j, (h, w) in enumerate(zip(gh, gw), start=1):
    cell = g.cell(row=3, column=j, value=h); cell.font = font(10, True, "FFFFFF"); cell.fill = fill(HDR); cell.border = border
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    g.column_dimensions[get_column_letter(j)].width = w
for k, row in enumerate(R["sector_guide"], start=4):
    for j, v in enumerate(row, start=1):
        cell = g.cell(row=k, column=j, value=v); cell.font = font(10); cell.border = border
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    a = g.cell(row=k, column=1); a.font = font(10, True, "FFFFFF"); a.fill = fill(SECTOR_COLOR.get(row[0], HDR))
    g.cell(row=k, column=7, value=f'=COUNTIF({rng("major")},A{k})')
    g.cell(row=k, column=8, value=f'=COUNTIFS({rng("major")},A{k},{rng("grade")},"A")')
    g.cell(row=k, column=9, value=f'=IFERROR(AVERAGEIFS({rng("upside")},{rng("major")},A{k}),"")')
    for j in (7, 8, 9):
        cc = g.cell(row=k, column=j); cc.font = font(10); cc.border = border; cc.alignment = Alignment(horizontal="center", vertical="top")
    g.cell(row=k, column=9).number_format = '0.0%'
gl_ = 4 + len(R["sector_guide"])
g.cell(row=gl_, column=1, value="합계").font = font(10, True)
g.cell(row=gl_, column=7, value=f"=SUM(G4:G{gl_-1})").font = font(10, True)
g.cell(row=gl_, column=8, value=f"=SUM(H4:H{gl_-1})").font = font(10, True)
g.freeze_panes = "B4"

# ======================= 요약 =======================
s = wb.create_sheet("요약")
s["A1"] = "시나리오·등급별 요약"; s["A1"].font = font(13, True, HDR)
sh = ["시나리오","종목 수","A등급","B등급","C등급","평균 목표가 상승여력","평균 모닝스타 할인율(수치 있는 종목)"]
for j, h in enumerate(sh, start=1):
    cell = s.cell(row=3, column=j, value=h); cell.font = font(10, True, "FFFFFF"); cell.fill = fill(HDR); cell.border = border
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
scens = ["지금 유리", "방어", "인상 종료 후", "AI 성장"]
for k, sc in enumerate(scens, start=4):
    s.cell(row=k, column=1, value=sc)
    s.cell(row=k, column=2, value=f'=COUNTIF({rng("scen")},A{k})')
    for jj, gr in enumerate("ABC", start=3):
        s.cell(row=k, column=jj, value=f'=COUNTIFS({rng("scen")},$A{k},{rng("grade")},"{gr}")')
    s.cell(row=k, column=6, value=f'=IFERROR(AVERAGEIFS({rng("upside")},{rng("scen")},A{k}),"")')
    s.cell(row=k, column=7, value=f'=IFERROR(AVERAGEIFS({rng("msdisc")},{rng("scen")},A{k}),"")')
tr = 4 + len(scens)
s.cell(row=tr, column=1, value="합계")
for jj in range(2, 6):
    Lc = get_column_letter(jj); s.cell(row=tr, column=jj, value=f"=SUM({Lc}4:{Lc}{tr-1})")
s.cell(row=tr, column=6, value=f'=IFERROR(AVERAGE({rng("upside")}),"")')
s.cell(row=tr, column=7, value=f'=IFERROR(AVERAGE({rng("msdisc")}),"")')
for rr in range(4, tr + 1):
    for jj in range(1, 8):
        cell = s.cell(row=rr, column=jj); cell.border = border; cell.font = font(10, rr == tr)
        if jj >= 6: cell.number_format = '0.0%'
s.column_dimensions["A"].width = 16
for Lc in "BCDE": s.column_dimensions[Lc].width = 9
s.column_dimensions["F"].width = 16; s.column_dimensions["G"].width = 22; s.row_dimensions[3].height = 32

# ======================= 열 설명·기준 =======================
cr = wb.create_sheet("열 설명·기준")
lines = [
("데이터 출처와 갱신", True),
(f"시세·시가총액·PER·예상 PER·배당·애널리스트 의견·목표가·52주 범위: StockAnalysis.com 종목 페이지를 {RUN_DATE}에 조회한 값입니다. 종목별 기준일은 '시세 기준일' 열에 있습니다.", False),
("현재가: 날짜가 찍힌 일봉 종가가 시세 페이지 값보다 같거나 최신이면 일봉 종가를 씁니다. 개인용 웹 앱에서는 일봉 차트(캔들)와 전일 대비 등락도 볼 수 있습니다.", False),
("이 파일은 매주 자동으로 다시 만들어집니다. '지난 조회 대비'는 직전 갱신 때 가격과 비교한 변화율입니다(첫 갱신 때는 '첫 조회').", False),
("모닝스타 공정가치·할인율·해자·불확실성: 모닝스타 공개 기사 기준입니다. 공개되지 않은 종목은 '비공개'로 표시했습니다. 모닝스타 수치와 종합등급은 매주 자동으로 바뀌지 않으며, 리서치를 다시 요청할 때 갱신됩니다.", False),
("", False),
("시트 안내", True),
("한눈에 보기: 섹터별로 묶은 핵심 지표 표 / 종목 카드: 회사별 카드 / 상세 데이터: 모든 열과 필터 / 섹터 가이드: 섹터별 사업·금리 영향 / 요약: 시나리오·등급별 집계", False),
("티커 칸의 색 = 섹터 색입니다. 티커를 누르면 회사 상세 페이지(사업 개요)가 열립니다.", False),
("", False),
("계산 열", True),
("지난 조회 대비 = 현재가 ÷ 지난 조회 가격 − 1", False),
("52주 고점 대비 = 현재가 ÷ 52주 최고가 − 1", False),
("목표가 대비 상승여력 = 애널리스트 평균 목표가 ÷ 현재가 − 1", False),
("모닝스타 기준 할인율 = 1 − 현재가 ÷ 모닝스타 공정가치 (공정가치가 없고 할인율만 발표된 경우 그 값을 사용)", False),
("파란 글씨(상세 데이터 시트) = 출처에서 가져온 입력값. 고치면 모든 시트의 계산이 자동으로 바뀝니다.", False),
("", False),
("종합등급 기준(Claude가 정리한 스크리닝 등급)", True),
("A: 저평가 근거가 뚜렷하고 이익 전망이 양호하며 재무·해자가 안정적(4가지 기준 중 3개 이상)", False),
("B: 저평가 근거는 있으나 뚜렷한 리스크가 1개 이상", False),
("C: 근거가 약하거나 이익 전망 하향·가치함정 위험, 또는 애널리스트 목표가가 현재가 수준 이하", False),
("4가지 기준: 1) 실적과 주가의 괴리  2) 밸류에이션  3) 이익 전망의 방향  4) 재무 안전성과 가까운 촉매", False),
("", False),
("시나리오", True),
("지금 유리(금리 인상 지속 시 유리) / 방어(금리와 무관한 방어력) / 인상 종료 후(금리 인하 전환 시 수혜) / AI 성장(AI 투자 사이클 연동)", False),
("", False),
("면책", True),
("공개 자료를 정리한 참고용 스크리닝 결과이며 투자 권유나 금융 자문이 아닙니다. 투자 결정 전 최신 시세와 공시를 확인하세요.", False),
]
for k, (txt, bold) in enumerate(lines, start=1):
    cell = cr.cell(row=k, column=1, value=txt)
    cell.font = font(11 if bold else 10, bold); cell.alignment = Alignment(wrap_text=True, vertical="top")
cr.column_dimensions["A"].width = 130

# ======================= 출처 =======================
so = wb.create_sheet("출처")
so["A1"] = "출처"; so["B1"] = "URL"
for cc in ("A1", "B1"): so[cc].font = font(10, True, "FFFFFF"); so[cc].fill = fill(HDR)
for k, (t, u) in enumerate(R["sources"], start=2):
    so.cell(row=k, column=1, value=t).font = font(10)
    cell = so.cell(row=k, column=2, value=u); cell.hyperlink = u; cell.font = font(10, color="0563C1", underline="single")
so.column_dimensions["A"].width = 70; so.column_dimensions["B"].width = 100

out = f"저평가_우량주_리스트_{RUN_DATE}.xlsx"
wb.save(out)

# ======================= 대시보드용 data.json =======================
def pct(a, b): return (a / b - 1) if (a is not None and b) else None
dash = []
for d in view_rows:
    msd = (1 - d["price"] / d["fv"]) if d["fv"] else d["disc"]
    dash.append({k: d.get(k) for k in ("t","name","major","sub","desc","rev","hq","scen","grade","price","qdate","mcap",
                                       "pe","fpe","div","rating","tgt","fv","moat","unc","val","thesis","risk","mdate","lo","hi")}
                | {"upside": pct(d["tgt"], d["price"]), "msdisc": msd, "off52": pct(d["price"], d["hi"]),
                   "chg": pct(d["price"], d["prev"]), "prevdate": d["prevdate"], "color": SECTOR_COLOR[d["major"]],
                   "dchg": pct(d["price"], d["ref_close"]), "diff": (d["price"] - d["ref_close"]) if d["ref_close"] else None, "ref_date": d["ref_date"], "src_price": d["src_price"],
                   "note": d.get("note"), "h": d["h"]})
# 시장 지수(코스피·나스닥 등): indices.json → 최근 150거래일
INDEX_META = [  # 기호, 이름, 짧은 이름, 지역, 단위, 실시간 링크
    ("^KS11", "코스피", "코스피", "국내", "pt", "https://finance.yahoo.com/quote/%5EKS11/"),
    ("^KQ11", "코스닥", "코스닥", "국내", "pt", "https://finance.yahoo.com/quote/%5EKQ11/"),
    ("^IXIC", "나스닥 종합", "나스닥", "미국", "pt", "https://finance.yahoo.com/quote/%5EIXIC/"),
    ("^GSPC", "S&P 500", "S&P 500", "미국", "pt", "https://finance.yahoo.com/quote/%5EGSPC/"),
    ("^DJI", "다우존스", "다우", "미국", "pt", "https://finance.yahoo.com/quote/%5EDJI/"),
    ("^SOX", "필라델피아 반도체", "반도체(SOX)", "미국", "pt", "https://finance.yahoo.com/quote/%5ESOX/"),
    ("KRW=X", "달러/원 환율", "달러 환율", "환율", "원", "https://finance.yahoo.com/quote/KRW%3DX/"),
    ("^TNX", "미국 10년물 국채 금리", "美 10년물", "금리", "%", "https://finance.yahoo.com/quote/%5ETNX/"),
    ("^VIX", "VIX 변동성지수", "VIX", "변동성", "pt", "https://finance.yahoo.com/quote/%5EVIX/"),
]
IDX = json.load(open("indices.json")) if os.path.exists("indices.json") else {}
indices = []
for sym, nm, short, region, unit, url in INDEX_META:
    h = IDX.get(sym, [])[-150:]
    if len(h) < 2: continue
    indices.append({"sym": sym, "name": nm, "short": short, "region": region, "unit": unit, "url": url, "h": h,
                    "price": h[-1][4], "qdate": iso(h[-1][0]), "ref_close": h[-2][4], "ref_date": iso(h[-2][0]),
                    "dchg": h[-1][4] / h[-2][4] - 1, "diff": h[-1][4] - h[-2][4]})

json.dump({"updated": RUN_DATE, "quote_date": max(d["qdate"] for d in rows), "stocks": dash, "indices": indices,
           "sector_guide": R["sector_guide"], "sector_colors": SECTOR_COLOR},
          open("data.json", "w", encoding="utf-8"), ensure_ascii=False)
print(out, len(rows), "종목")

# ======================= 대시보드 HTML =======================
if os.path.exists("dashboard_template.html"):
    html = open("dashboard_template.html", encoding="utf-8").read()
    payload = json.dumps(json.load(open("data.json", encoding="utf-8")), ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    open("dashboard.html", "w", encoding="utf-8").write(html.replace("__DATA__", payload))
    print("dashboard.html 생성")
